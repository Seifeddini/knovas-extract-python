"""DOCX layout mode: tables rendered in place with the PDF row grammar."""

from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape

import pytest

pytest.importorskip("docx")
import docx as _docx  # noqa: E402

from knovas_extract import extract  # noqa: E402
from knovas_extract._layout.lint import est_tokens  # noqa: E402
from knovas_extract.extractors.docx import DocxExtractor  # noqa: E402
from knovas_extract.result import Limits  # noqa: E402

pytestmark = [pytest.mark.unit]

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ROW_MAX_TOKENS, ROW_MAX_CHARS = 200, 1800  # LayoutOptions defaults (docs/layout-text-mode.md)
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _save(d: object) -> bytes:
    buf = io.BytesIO()
    d.save(buf)  # type: ignore[attr-defined]
    return buf.getvalue()


def _with_body(body_xml: str) -> bytes:
    """A DOCX whose body is ``body_xml`` (shapes python-docx cannot write: a row that
    starts late, a vertical-merge continuation without a cell above)."""
    src = zipfile.ZipFile(io.BytesIO(_save(_docx.Document())))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "word/document.xml":
                data = (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    f'<w:document xmlns:w="{_W}"><w:body>{body_xml}<w:sectPr/></w:body></w:document>'
                ).encode()
            z.writestr(info.filename, data)
    return out.getvalue()


def _p(text: str = "") -> str:
    run = f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r>' if text else ""
    return f"<w:p>{run}</w:p>"


def _tc(text: str = "", *, span: int = 1, vmerge: str | None = None) -> str:
    pr = f'<w:gridSpan w:val="{span}"/>' if span > 1 else ""
    pr += {None: "", "restart": '<w:vMerge w:val="restart"/>', "continue": "<w:vMerge/>"}[vmerge]
    return f"<w:tc><w:tcPr>{pr}</w:tcPr>{_p(text)}</w:tc>"


def _tr(*cells: str, before: int = 0) -> str:
    pr = f'<w:trPr><w:gridBefore w:val="{before}"/></w:trPr>' if before else ""
    return f"<w:tr>{pr}{''.join(cells)}</w:tr>"


def _tbl(*rows: str, ncols: int) -> str:
    grid = '<w:gridCol w:w="2000"/>' * ncols
    return f"<w:tbl><w:tblPr/><w:tblGrid>{grid}</w:tblGrid>{''.join(rows)}</w:tbl>"


def _words(n: int, stem: str) -> str:
    return " ".join(f"{stem}{i}" for i in range(n))


def _assert_lines_bounded(text: str) -> None:
    for line in text.split("\n"):
        assert len(line) <= ROW_MAX_CHARS, line[:60]
        assert est_tokens(line) <= ROW_MAX_TOKENS, line[:60]


def _doc_with_table() -> bytes:
    d = _docx.Document()
    d.add_paragraph("Prosa davor.")
    t = d.add_table(rows=3, cols=3)
    for c, h in enumerate(["Position", "2023", "2022"]):
        t.rows[0].cells[c].text = h
    for c, v in enumerate(["Umsatz", "1'234.00", "1'100.00"]):
        t.rows[1].cells[c].text = v
    merged = t.rows[2].cells[0].merge(t.rows[2].cells[1])
    merged.text = "Total"
    t.rows[2].cells[2].text = "9'999.00"
    d.add_paragraph("Prosa danach.")
    return _save(d)


def test_without_tables_layout_is_byte_identical_to_plain() -> None:
    d = _docx.Document()
    d.add_heading("Vertrag", level=1)
    d.add_paragraph("Die Parteien vereinbaren Folgendes.")
    data = _save(d)
    plain = extract(data, mime=DOCX)
    layout = extract(data, mime=DOCX, text_mode="layout")
    assert layout.content.text == plain.content.text
    assert layout.content.sections == plain.content.sections
    assert layout.metadata.extra["docx:text_mode"] == "layout"
    assert layout.metadata.extra["docx:layout_tables"] == 0
    assert "docx:text_mode" not in plain.metadata.extra
    assert not any("text_mode=" in w for w in layout.warnings)


def test_table_is_rendered_in_place_with_fold_keys() -> None:
    r = extract(_doc_with_table(), mime=DOCX, text_mode="layout")
    assert r.content.text == (
        "Prosa davor.\n\n"
        "Position | 2023 | 2022\n"
        "Umsatz | 2023: 1'234.00 | 2022: 1'100.00\n"
        "Total | 2022: 9'999.00\n\n"
        "Prosa danach."
    )
    assert r.metadata.extra["docx:layout_tables"] == 1


def test_plain_mode_and_content_tables_are_unchanged() -> None:
    data = _doc_with_table()
    plain = extract(data, mime=DOCX)
    layout = extract(data, mime=DOCX, text_mode="layout")
    assert plain.content.text == "Prosa davor.\n\nProsa danach."
    assert layout.content.tables == plain.content.tables


def test_header_repeats_on_every_pack() -> None:
    d = _docx.Document()
    t = d.add_table(rows=41, cols=3)
    for c, h in enumerate(["Konto", "2023", "2022"]):
        t.rows[0].cells[c].text = h
    for i in range(1, 41):
        for c, v in enumerate([f"Konto {i}", f"{i}'000.00", f"{i}'500.00"]):
            t.rows[i].cells[c].text = v
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    packs = text.split("\n\n")
    assert len(packs) > 1
    assert all(p.split("\n")[0] == "Konto | 2023 | 2022" for p in packs)
    assert sum(len(p.split("\n")) - 1 for p in packs) == 40


def test_nested_table_text_is_kept() -> None:
    d = _docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Feld"
    t.rows[0].cells[1].text = "Wert"
    t.rows[1].cells[0].text = "Adresse"
    inner = t.rows[1].cells[1].add_table(rows=1, cols=2)
    inner.rows[0].cells[0].text = "Bahnhofstrasse 1"
    inner.rows[0].cells[1].text = "8001 Zürich"
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    assert "Bahnhofstrasse 1" in text and "8001 Zürich" in text


def test_single_row_table_becomes_one_line() -> None:
    d = _docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Muster AG"
    t.rows[0].cells[1].text = "Basel"
    assert extract(_save(d), mime=DOCX, text_mode="layout").content.text == "Muster AG | Basel"


def test_sections_point_at_their_heading_line_in_the_layout_text() -> None:
    d = _docx.Document()
    t = d.add_table(rows=2, cols=2)
    for c, v in enumerate(["A", "B"]):
        t.rows[0].cells[c].text = v
    for c, v in enumerate(["x", "1.00"]):
        t.rows[1].cells[c].text = v
    d.add_heading("Anhang", level=1)
    d.add_paragraph("Text im Anhang.")
    r = extract(_save(d), mime=DOCX, text_mode="layout", emit_sentences=True)
    sec = next(s for s in r.content.sections if s.heading == "Anhang")
    assert r.content.text.split("\n")[sec.line_start - 1] == "Anhang"
    assert r.content.sentences  # dispatch's sentence contracts held


def test_layout_output_is_deterministic() -> None:
    data = _doc_with_table()
    a = extract(data, mime=DOCX, text_mode="layout")
    b = extract(data, mime=DOCX, text_mode="layout")
    assert a.content.text == b.content.text


def test_other_formats_warn_and_stay_plain() -> None:
    r = extract(b"Nur Text.", mime="text/plain", text_mode="layout")
    assert (
        "text_mode='layout' is implemented for PDF and DOCX only; plain text emitted" in r.warnings
    )


# --- row and pack bounds: DOCX cells, unlike a PDF line, have no length limit ---


def test_long_label_is_split_like_a_cell_not_repeated() -> None:
    # A Synopse (old law | new law): every cell is longer than the row budget.
    d = _docx.Document()
    t = d.add_table(rows=21, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Geltendes Recht", "Neues Recht"
    source = 0
    for i in range(1, 21):
        old, new = f"Art. {i} " + _words(150, "alt"), f"Art. {i} " + _words(150, "neu")
        t.rows[i].cells[0].text, t.rows[i].cells[1].text = old, new
        source += len(old) + len(new)
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    _assert_lines_bounded(text)
    assert text.count("alt149") == 20  # every label once: nothing lost, nothing repeated
    assert text.count("neu149") == 20
    assert len(text) <= 1.5 * source + 4096  # the PDF expansion invariant (g)


def test_long_label_only_rows_and_one_column_rows_are_split() -> None:
    d = _docx.Document()
    t = d.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Hinweis", "Betrag"
    t.rows[1].cells[0].text = _words(300, "wort")
    t.rows[2].cells[0].text, t.rows[2].cells[1].text = "Total", "100.00"
    d.add_paragraph("Zwischentext.")
    box = d.add_table(rows=2, cols=1)
    box.rows[0].cells[0].text = "Wichtiger Hinweis"
    box.rows[1].cells[0].text = _words(300, "box")
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    _assert_lines_bounded(text)
    assert "wort299" in text and "box299" in text and "(Forts.) | " in text
    assert "Total | Betrag: 100.00" in text


def test_long_header_is_emitted_once_not_on_every_pack() -> None:
    d = _docx.Document()
    t = d.add_table(rows=21, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = _words(300, "kopf"), "Wert"
    for i in range(1, 21):
        t.rows[i].cells[0].text, t.rows[i].cells[1].text = f"Zeile {i}", f"{i}.00"
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    _assert_lines_bounded(text)
    assert text.count("kopf299") == 1
    assert "Zeile 20 | Wert: 20.00" in text  # the fold keys stay


def test_one_row_table_over_the_row_limit_is_split() -> None:
    d = _docx.Document()
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = _words(300, "wort")
    text = extract(_save(d), mime=DOCX, text_mode="layout").content.text
    assert text.count("\n") > 0 and "wort299" in text
    _assert_lines_bounded(text)


def test_a_carriage_return_in_a_cell_keeps_the_row_on_one_line() -> None:
    # `&#13;` survives XML parsing; canonicalize_text would make it a line break.
    cell = (
        '<w:tc><w:p><w:r><w:t xml:space="preserve">Umsatz&#13;&#13;# netto</w:t></w:r></w:p></w:tc>'
    )
    data = _with_body(_tbl(_tr(_tc("Position"), _tc("2023")), _tr(cell, _tc("1.00")), ncols=2))
    assert extract(data, mime=DOCX, text_mode="layout").content.text == (
        "Position | 2023\nUmsatz # netto | 2023: 1.00"
    )


def test_a_token_longer_than_a_row_is_cut() -> None:
    blob = "QUJD" * 1500
    data = _with_body(_tbl(_tr(_tc("Feld"), _tc("Wert")), _tr(_tc("Signatur"), _tc(blob)), ncols=2))
    text = extract(data, mime=DOCX, text_mode="layout").content.text
    _assert_lines_bounded(text)
    assert text.count("Q") == 1500  # cut, not dropped


# --- fail-soft: what plain mode survives, layout mode survives ---


@pytest.mark.parametrize(
    ("rows", "expected"),
    [
        pytest.param(  # a continuation in the first row: no cell above it
            (
                _tr(_tc("A", vmerge="continue"), _tc("2023"), _tc("2022")),
                _tr(_tc("Umsatz"), _tc("1'234.00"), _tc("1'100.00")),
            ),
            "A | 2023 | 2022\nUmsatz | 2023: 1'234.00 | 2022: 1'100.00",
            id="continuation-in-first-row",
        ),
        pytest.param(  # a continuation under a row whose cell there spans two columns
            (
                _tr(_tc("Position"), _tc("2023"), _tc("2022")),
                _tr(_tc("Titel", span=2), _tc("x")),
                _tr(_tc("Umsatz"), _tc("1'234.00", vmerge="continue"), _tc("1'100.00")),
            ),
            "Position | 2023 | 2022\nTitel | 2022: x\nUmsatz | 2023: 1'234.00 | 2022: 1'100.00",
            id="continuation-under-a-span",
        ),
    ],
)
def test_vertical_merge_without_a_cell_above_renders_the_cell_itself(
    rows: tuple[str, ...], expected: str
) -> None:
    data = _with_body(_p("Vorher.") + _tbl(*rows, ncols=3) + _p("Nachher."))
    plain = extract(data, mime=DOCX)
    layout = extract(data, mime=DOCX, text_mode="layout")
    assert plain.content.text == "Vorher.\n\nNachher."  # python-docx gives up on this table
    assert layout.content.text == f"Vorher.\n\n{expected}\n\nNachher."
    assert layout.metadata.extra["docx:layout_tables"] == 1


def test_long_vertical_merge_is_read_without_recursion() -> None:
    from knovas_extract.extractors.docx import _grid_rows

    rows = [_tr(_tc("Gruppe"), _tc("Konto"), _tc("2023"))]
    rows.append(_tr(_tc("Kasse", vmerge="restart"), _tc("Konto 1"), _tc("1.00")))
    rows += [_tr(_tc(vmerge="continue"), _tc(f"Konto {i}"), _tc(f"{i}.00")) for i in range(2, 1501)]
    table = _docx.Document(io.BytesIO(_with_body(_tbl(*rows, ncols=3)))).tables[0]
    grid = _grid_rows(table)  # python-docx's row.cells: RecursionError, after seconds
    assert len(grid) == 1501
    assert grid[-1] == ["Kasse", "Konto 1500", "1500.00"]  # as content.tables repeats it


def test_a_table_that_fails_to_render_is_left_out_with_a_counted_warning(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from knovas_extract.extractors import docx as docx_mod

    def boom(table: object, budget: int) -> str:
        raise ValueError("no tr above topmost tr in w:tbl")

    monkeypatch.setattr(docx_mod, "_layout_table_block", boom)
    r = extract(_doc_with_table(), mime=DOCX, text_mode="layout")
    assert r.content.text == "Prosa davor.\n\nProsa danach."  # what plain mode emits
    assert r.metadata.extra["docx:layout_tables"] == 0
    assert "docx: layout rendering failed on 1 table (ValueError); omitted from text" in r.warnings
    assert not any("w:tbl" in w for w in r.warnings)  # never the exception text


def test_a_table_over_the_text_limit_is_left_out_not_raised() -> None:
    # Called below extract(): a DOCX small enough for a 5 kB `max_text_bytes`
    # does not exist (_guard_zip checks every part, the template's styles are 438 kB).
    from knovas_extract.extractors.docx import _extract_body_layout

    d = _docx.Document()
    d.add_paragraph("Vorher.")
    big = d.add_table(rows=31, cols=2)
    big.rows[0].cells[0].text, big.rows[0].cells[1].text = "Konto", "Notiz"
    for i in range(1, 31):
        big.rows[i].cells[0].text, big.rows[i].cells[1].text = f"Konto {i}", _words(100, "n")
    d.add_paragraph("Nachher.")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Klein"  # still fits after it
    text, tables, warnings = _extract_body_layout(_save(d), Limits(max_text_bytes=5_000))
    assert text == "Vorher.\n\nNachher.\n\nKlein"
    assert tables == ["Klein"]
    assert warnings == ["docx: 1 table over the text size limit omitted from text"]


def test_vertical_merge_repeats_a_short_label_and_writes_a_long_note_once() -> None:
    note = _words(40, "notiz")  # over a quarter of the row budget
    data = _with_body(
        _tbl(
            _tr(_tc("Gruppe"), _tc("Konto"), _tc("2023"), _tc("Notiz")),
            _tr(
                _tc("Umlaufvermögen", vmerge="restart"),
                _tc("Kasse"),
                _tc("1.00"),
                _tc(note, vmerge="restart"),
            ),
            _tr(_tc(vmerge="continue"), _tc("Bank"), _tc("2.00"), _tc(vmerge="continue")),
            ncols=4,
        )
    )
    assert extract(data, mime=DOCX, text_mode="layout").content.text == (
        "Gruppe | Konto | 2023 | Notiz\n"
        f"Umlaufvermögen | Kasse | 2023: 1.00 | {note}\n\n"
        "Gruppe | Konto | 2023 | Notiz\n"
        "Umlaufvermögen | Bank | 2023: 2.00"
    )


# --- grid alignment and sections ---


def test_rows_that_start_late_keep_their_column_keys() -> None:
    data = _with_body(
        _tbl(
            _tr(_tc("Position"), _tc("Anhang"), _tc("2023"), _tc("2022")),
            _tr(_tc("Umsatz"), _tc("2.1"), _tc("1'234.00"), _tc("1'100.00")),
            _tr(_tc("-300.00"), _tc("-250.00"), before=2),
            _tr(_tc("777.00"), before=3),
            ncols=4,
        )
    )
    assert extract(data, mime=DOCX, text_mode="layout").content.text == (
        "Position | Anhang | 2023 | 2022\n"
        "Umsatz | Anhang: 2.1 | 2023: 1'234.00 | 2022: 1'100.00\n"
        "2023: -300.00 | 2022: -250.00\n"
        "2022: 777.00"
    )


def test_sections_skip_table_text_that_repeats_a_heading() -> None:
    d = _docx.Document()
    d.add_heading("Jahresrechnung 2023", level=1)
    d.add_paragraph("Bilanz per 31. Dezember 2023")
    t = d.add_table(rows=3, cols=4)
    rows = [("Aktiven", "Anhang", "2023", "2022"), ("Kasse", "2.1", "1.00", "2.00")]
    rows.append(("Bank", "2.2", "3.00", "4.00"))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    d.add_heading("Anhang", level=1)
    d.add_paragraph("Die Jahresrechnung folgt dem OR.")
    r = extract(_save(d), mime=DOCX, text_mode="layout", emit_sentences=True)
    lines = r.content.text.split("\n")
    first, anhang = r.content.sections
    assert lines[first.line_start - 1] == "Jahresrechnung 2023"
    assert lines[anhang.line_start - 1] == "Anhang"
    assert first.line_end is not None and first.line_end < anhang.line_start
    rows_ = [s for s in r.content.sentences if " | " in s.text]
    assert rows_ and all(s.section_index == 0 for s in rows_)


def test_sections_skip_an_overview_table_that_lists_the_headings() -> None:
    d = _docx.Document()
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Kapitel", "Seite"), ("Einleitung", "1"), ("Anhang", "2")]):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    d.add_heading("Einleitung", level=1)
    d.add_paragraph("Text der Einleitung.")
    d.add_heading("Anhang", level=1)
    d.add_paragraph("Text im Anhang.")
    r = extract(_save(d), mime=DOCX, text_mode="layout", emit_sentences=True)
    lines = r.content.text.split("\n")
    assert [lines[s.line_start - 1] for s in r.content.sections] == ["Einleitung", "Anhang"]
    by_text = {s.text: s.section_index for s in r.content.sentences}
    assert by_text["Text der Einleitung."] == 0 and by_text["Text im Anhang."] == 1


def test_docx_extractor_rejects_an_unknown_text_mode() -> None:
    with pytest.raises(ValueError, match="text_mode"):
        DocxExtractor().extract(_doc_with_table(), text_mode="Layout")  # type: ignore[arg-type]
